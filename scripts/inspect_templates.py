import json, openpyxl
from docx import Document
w=openpyxl.load_workbook('references/Day22_template.xlsx')
out={}
for s in w:
    lines=[]
    for row in s:
        cells=[]
        for c in row:
            if c.value is not None or (c.has_style and c.fill.fgColor.rgb in ['00FFF2CC','00FFFF00','FFFFFF00','FFFFF2CC']):
                cells.append({'cell':c.coordinate,'value':c.value,'fill':str(c.fill.fgColor.rgb),'style':c.style_id})
        if cells: lines.append(cells)
    out[s.title]={'merges':[str(x) for x in s.merged_cells.ranges],'rows':lines}
open('references/template_cells.json','w',encoding='utf-8').write(json.dumps(out,ensure_ascii=False,indent=2,default=str))
d=Document('references/Day22_template.docx')
text=[]
for p in d.paragraphs: text.append(p.text)
for t in d.tables:
    for r in t.rows: text.append(' | '.join(c.text for c in r.cells))
open('references/template_doc.txt','w',encoding='utf-8').write('\n'.join(text))
print('Template inspection saved')
